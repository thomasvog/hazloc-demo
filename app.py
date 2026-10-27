"""Inspecciones HAZLOC (IEC/IECEx) - versión para Vercel + Postgres (Neon)."""
import csv
import hashlib
import io
import os
import re
import secrets
import unicodedata
from datetime import date
from functools import wraps

import psycopg
from flask import (Flask, abort, flash, g, redirect, render_template, request,
                   send_file, session, url_for)
from jinja2 import DictLoader
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageOps
from psycopg.rows import dict_row
from werkzeug.security import check_password_hash, generate_password_hash

from tpl import TEMPLATES

URL = os.environ.get("DATABASE_URL") or os.environ.get("POSTGRES_URL") or ""
SETUP_KEY = os.environ.get("SETUP_KEY", "")

app = Flask(__name__)
app.jinja_env.loader = DictLoader(TEMPLATES)
app.config.update(
    SECRET_KEY=os.environ.get("SECRET_KEY") or hashlib.sha256(("hazloc|" + URL).encode()).hexdigest(),
    MAX_CONTENT_LENGTH=4 * 1024 * 1024,  # Vercel admite ~4,5 MB por request
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=bool(os.environ.get("VERCEL")),
)

AREAS = ["Polímeros", "Extrusión", "Unidad 500", "Servicios", "Bagging"]  # áreas fijas
ZONES = ["0", "1", "2", "20", "21", "22", "0/20", "1/21", "2/22", "Segura"]
ZMAP = {z.lower().replace(" ", ""): z for z in ZONES}
ZMAP.update({"noclasificada": "Segura", "safe": "Segura"})
AREA_COLORS = ["#2563a8", "#7a4fb5", "#1aa3c8", "#5b6b7b", "#c2409b"]  # azul, violeta, celeste, gris, magenta
CRIT = ["alta", "media", "baja"]
PROT = ["Ex d", "Ex e", "Ex i", "Ex p", "Ex n", "Ex m", "Ex o", "Ex q", "Ex t", "Otro / N.A."]
ROLES = {"tecnico": 1, "supervisor": 2, "admin": 3}
ROLE_LABEL = {"tecnico": "Técnico", "supervisor": "Supervisor", "admin": "Administrador"}

SCHEMA = [
    "CREATE TABLE IF NOT EXISTS users(id SERIAL PRIMARY KEY, username TEXT UNIQUE NOT NULL, name TEXT NOT NULL, pw TEXT NOT NULL, role TEXT NOT NULL)",
    "CREATE TABLE IF NOT EXISTS locations(id SERIAL PRIMARY KEY, parent_id INT REFERENCES locations(id), name TEXT NOT NULL)",
    """CREATE TABLE IF NOT EXISTS findings(id SERIAL PRIMARY KEY, location_id INT NOT NULL REFERENCES locations(id),
       tag TEXT, zone TEXT NOT NULL, protection TEXT, criticality TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'pendiente',
       description TEXT NOT NULL, inspector_id INT REFERENCES users(id) ON DELETE SET NULL, inspected_at DATE NOT NULL,
       resolution TEXT, closed_by INT REFERENCES users(id) ON DELETE SET NULL, closed_at DATE)""",
    """CREATE TABLE IF NOT EXISTS photos(id SERIAL PRIMARY KEY, finding_id INT NOT NULL REFERENCES findings(id) ON DELETE CASCADE,
       kind TEXT NOT NULL, data BYTEA NOT NULL, created_at DATE NOT NULL)""",
    "CREATE TABLE IF NOT EXISTS attempts(id SERIAL PRIMARY KEY, username TEXT NOT NULL, at TIMESTAMPTZ NOT NULL DEFAULT now())",
    "ALTER TABLE users ADD COLUMN IF NOT EXISTS active BOOLEAN NOT NULL DEFAULT TRUE",
    """CREATE TABLE IF NOT EXISTS equipment(id SERIAL PRIMARY KEY, tag TEXT NOT NULL, description TEXT,
       location_id INT NOT NULL REFERENCES locations(id), zone TEXT NOT NULL, protection TEXT)""",
    "CREATE UNIQUE INDEX IF NOT EXISTS ux_eq_tag ON equipment(lower(tag))",
    "ALTER TABLE findings ADD COLUMN IF NOT EXISTS equipment_id INT REFERENCES equipment(id) ON DELETE SET NULL",
    "ALTER TABLE findings ADD COLUMN IF NOT EXISTS inspector_name TEXT",
    "ALTER TABLE findings ADD COLUMN IF NOT EXISTS closer_name TEXT",
    "ALTER TABLE findings ALTER COLUMN inspector_id DROP NOT NULL",
    "UPDATE findings f SET inspector_name=u.name FROM users u WHERE u.id=f.inspector_id AND f.inspector_name IS NULL",
    "UPDATE findings f SET closer_name=u.name FROM users u WHERE u.id=f.closed_by AND f.closer_name IS NULL",
    *[("DO $$ BEGIN IF EXISTS (SELECT 1 FROM pg_constraint WHERE conname='findings_%s_fkey' AND confdeltype <> 'n') THEN "
       "ALTER TABLE findings DROP CONSTRAINT findings_%s_fkey; "
       "ALTER TABLE findings ADD CONSTRAINT findings_%s_fkey FOREIGN KEY (%s) REFERENCES users(id) ON DELETE SET NULL; "
       "END IF; END $$") % (c, c, c, c) for c in ("inspector_id", "closed_by")],
    """CREATE TABLE IF NOT EXISTS declassifications(id SERIAL PRIMARY KEY, location_id INT NOT NULL REFERENCES locations(id),
       reason TEXT NOT NULL, description TEXT NOT NULL, done_at DATE NOT NULL, responsible TEXT NOT NULL,
       created_by INT REFERENCES users(id) ON DELETE SET NULL, set_safe BOOLEAN NOT NULL DEFAULT FALSE)""",
    """CREATE TABLE IF NOT EXISTS declass_tags(id SERIAL PRIMARY KEY, declass_id INT NOT NULL REFERENCES declassifications(id) ON DELETE CASCADE,
       equipment_id INT REFERENCES equipment(id) ON DELETE SET NULL, tag TEXT NOT NULL, prev_zone TEXT)""",
    """CREATE TABLE IF NOT EXISTS declass_photos(id SERIAL PRIMARY KEY, declass_id INT NOT NULL REFERENCES declassifications(id) ON DELETE CASCADE,
       kind TEXT NOT NULL, data BYTEA NOT NULL, created_at DATE NOT NULL)""",
    "CREATE INDEX IF NOT EXISTS ix_dt ON declass_tags(declass_id)",
    "CREATE INDEX IF NOT EXISTS ix_dp ON declass_photos(declass_id)",
    "UPDATE users SET role='tecnico' WHERE role='carga'",
    "UPDATE users SET role='supervisor' WHERE role='cierre'",
    "ALTER TABLE equipment ALTER COLUMN location_id DROP NOT NULL",
    "ALTER TABLE equipment ALTER COLUMN zone DROP NOT NULL",
    "CREATE INDEX IF NOT EXISTS ix_f ON findings(status, criticality, zone, location_id)",
    "CREATE INDEX IF NOT EXISTS ix_p ON photos(finding_id)",
]
_ready = _has = False


def sync_areas():
    """Deja en la base exactamente las áreas fijas. Lo que esté en otras áreas (datos de prueba anteriores) se elimina."""
    run("UPDATE locations SET parent_id=NULL WHERE parent_id IS NOT NULL")
    keep = {}
    for r in rows("SELECT id, name FROM locations ORDER BY id"):
        if r["name"] in AREAS and r["name"] not in keep:
            keep[r["name"]] = r["id"]
        elif r["name"] in AREAS:  # duplicado de un área fija: se reasigna a la primera
            run("UPDATE findings SET location_id=%s WHERE location_id=%s", keep[r["name"]], r["id"])
            run("UPDATE equipment SET location_id=%s WHERE location_id=%s", keep[r["name"]], r["id"])
            run("DELETE FROM locations WHERE id=%s", r["id"])
    run("DELETE FROM findings WHERE location_id IN (SELECT id FROM locations WHERE name <> ALL(%s))", AREAS)
    run("DELETE FROM equipment WHERE location_id IN (SELECT id FROM locations WHERE name <> ALL(%s))", AREAS)
    run("DELETE FROM locations WHERE name <> ALL(%s)", AREAS)
    try:
        run("CREATE UNIQUE INDEX IF NOT EXISTS ux_loc_name ON locations(name)")
    except psycopg.Error:
        pass
    for n in AREAS:
        run("INSERT INTO locations(name) SELECT %s WHERE NOT EXISTS (SELECT 1 FROM locations WHERE name=%s)", n, n)


def db():
    if "db" not in g:
        g.db = psycopg.connect(URL, row_factory=dict_row, autocommit=True, prepare_threshold=None)
    return g.db


@app.teardown_appcontext
def close_db(_):
    d = g.pop("db", None)
    if d:
        d.close()


def rows(sql, *p):
    return db().execute(sql, p).fetchall()


def one(sql, *p):
    r = db().execute(sql, p).fetchone()
    if r is None:
        abort(404)
    return r


def run(sql, *p):
    return db().execute(sql, p)


def dt(s):
    try:
        return date.fromisoformat(s)
    except (TypeError, ValueError):
        return date.today()


def has_users():
    global _has
    if not _has:
        _has = db().execute("SELECT 1 FROM users LIMIT 1").fetchone() is not None
    return _has


def need(role="tecnico"):
    def deco(fn):
        @wraps(fn)
        def w(*a, **k):
            if "uid" not in session:
                return redirect(url_for("login", next=request.path))
            if ROLES[session["role"]] < ROLES[role]:
                abort(403)
            return fn(*a, **k)
        return w
    return deco


@app.before_request
def gate():
    global _ready
    if not URL:
        return "Falta conectar la base de datos (variable DATABASE_URL). Ver el README, paso 3.", 503
    if not _ready:
        for s in SCHEMA:
            try:
                db().execute(s)
            except psycopg.Error:
                pass
        try:
            sync_areas()
        except psycopg.Error:
            pass
        _ready = True
    session.setdefault("csrf", secrets.token_hex(16))
    if request.method == "POST" and request.form.get("_csrf") != session["csrf"]:
        abort(400)
    if request.endpoint not in ("setup", None) and not has_users():
        return redirect(url_for("setup"))
    if "uid" in session:
        u = db().execute("SELECT role, active, name FROM users WHERE id=%s", (session["uid"],)).fetchone()
        if not u or not u["active"]:
            session.clear()
            return redirect(url_for("login"))
        session["role"], session["name"] = u["role"], u["name"]


@app.errorhandler(413)
def too_big(_):
    return "Las fotos pesan demasiado (máximo 4 MB por envío). Volvé atrás y subí menos fotos a la vez.", 413


@app.context_processor
def consts():
    return dict(ZONES=ZONES, CRIT=CRIT, PROT=PROT, ROLES=ROLES, ROLE_LABEL=ROLE_LABEL,
                can_tags=ROLES.get(session.get("role"), 0) >= 2)


def locs():
    r = rows("SELECT id, name FROM locations ORDER BY array_position(%s::text[], name), name", AREAS)
    return [(x["id"], x["name"], 1) for x in r], {x["id"]: x["name"] for x in r}


def findings():
    a, w, p = request.args, [], []
    if a.get("loc", "").isdigit():
        w.append("f.location_id=%s")
        p.append(int(a["loc"]))
    if a.get("overdue"):
        w.append("f.status='pendiente' AND CURRENT_DATE - f.inspected_at > 30")
    for k, col in (("zone", "zone"), ("crit", "criticality"), ("status", "status")):
        if a.get(k):
            w.append("f.%s=%%s" % col)
            p.append(a[k])
    if a.get("q"):
        w.append("(f.tag ILIKE %s OR f.description ILIKE %s OR f.resolution ILIKE %s)")
        p += ["%" + a["q"] + "%"] * 3
    sql = ("SELECT f.*, COALESCE(u.name, f.inspector_name, 'Usuario eliminado') inspector, "
           "CASE WHEN f.closed_at IS NOT NULL THEN COALESCE(c.name, f.closer_name, 'Usuario eliminado') END closer FROM findings f "
           "LEFT JOIN users u ON u.id=f.inspector_id LEFT JOIN users c ON c.id=f.closed_by"
           + (" WHERE " + " AND ".join(w) if w else "")
           + " ORDER BY f.status='solucionado', CASE f.criticality WHEN 'alta' THEN 0 WHEN 'media' THEN 1 ELSE 2 END,"
             " f.inspected_at DESC, f.id DESC")
    return rows(sql, *p)


def jpeg_under(im, limit=300 * 1024):
    """Guarda la imagen como JPEG de hasta `limit` bytes bajando calidad y, si hace falta, tamaño."""
    q = 82
    while True:
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=q, optimize=True)
        if buf.tell() <= limit or min(im.size) <= 320:
            return buf.getvalue()
        if q > 45:
            q -= 8
        else:
            im = im.resize((int(im.width * .85), int(im.height * .85)), Image.LANCZOS)


def add_photos(fid, kind, tbl="photos", col="finding_id", field="photos"):
    n = 0
    for f in request.files.getlist(field):
        if not f.filename:
            continue
        try:
            im = ImageOps.exif_transpose(Image.open(f.stream)).convert("RGB")
            im.thumbnail((1600, 1600))
        except Exception:
            flash("No se pudo leer «%s»: no es una imagen válida." % f.filename)
            continue
        buf = io.BytesIO()
        buf.write(jpeg_under(im))
        run("INSERT INTO %s(%s,kind,data,created_at) VALUES(%%s,%%s,%%s,%%s)" % (tbl, col),
            fid, kind, buf.getvalue(), date.today())
        n += 1
    return n


@app.route("/setup", methods=["GET", "POST"])
def setup():
    global _has
    if has_users():
        return redirect(url_for("login"))
    if request.method == "POST":
        f = request.form
        un, nm, pw = f.get("username", "").strip().lower(), f.get("name", "").strip(), f.get("password", "")
        if not SETUP_KEY or not secrets.compare_digest(f.get("key", ""), SETUP_KEY):
            flash("Código de instalación incorrecto (es el valor de la variable SETUP_KEY en Vercel).")
        elif not un or not nm or len(pw) < 8:
            flash("Completá usuario y nombre. La contraseña debe tener al menos 8 caracteres.")
        else:
            run("INSERT INTO users(username,name,pw,role) VALUES(%s,%s,%s,'admin')", un, nm, generate_password_hash(pw))
            _has = True
            flash("Administrador creado. Ya podés ingresar.")
            return redirect(url_for("login"))
    return render_template("setup.html", key_missing=not SETUP_KEY)


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        un = request.form.get("username", "").strip().lower()
        if one("SELECT COUNT(*) n FROM attempts WHERE username=%s AND at > now() - interval '10 minutes'", un)["n"] >= 5:
            flash("Demasiados intentos fallidos. Esperá 10 minutos.")
            return render_template("login.html")
        u = db().execute("SELECT * FROM users WHERE username=%s", (un,)).fetchone()
        if u and u['active'] and check_password_hash(u["pw"], request.form.get("password", "")):
            run("DELETE FROM attempts WHERE username=%s", un)
            session.clear()
            session.update(uid=u["id"], name=u["name"], role=u["role"], csrf=secrets.token_hex(16))
            nxt = request.args.get("next", "/")
            return redirect(nxt if nxt.startswith("/") and not nxt.startswith("//") else "/")
        run("INSERT INTO attempts(username) VALUES(%s)", un)
        flash("Usuario o contraseña incorrectos.")
    return render_template("login.html")


@app.route("/logout", methods=["POST"])
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/")
@need()
def dashboard():
    st = {r["status"]: r["n"] for r in rows("SELECT status, COUNT(*) n FROM findings GROUP BY 1")}
    by_crit = {r["k"]: r["n"] for r in rows("SELECT criticality k, COUNT(*) n FROM findings WHERE status='pendiente' GROUP BY 1")}
    cnt = {r["location_id"]: r["n"] for r in rows("SELECT location_id, COUNT(*) n FROM findings WHERE status='pendiente' GROUP BY 1")}
    paths = locs()[1]
    by_area = [(i, n, AREA_COLORS[AREAS.index(n)] if n in AREAS else "#888888", cnt.get(i, 0)) for i, n in paths.items()]
    old = one("SELECT COUNT(*) n FROM findings WHERE status='pendiente' AND CURRENT_DATE - inspected_at > 30")["n"]
    top = rows("SELECT id, tag, zone, criticality, location_id, inspected_at, description FROM findings "
               "WHERE status='pendiente' ORDER BY CASE criticality WHEN 'alta' THEN 0 WHEN 'media' THEN 1 ELSE 2 END, inspected_at LIMIT 8")
    return render_template("dashboard.html", pend=st.get("pendiente", 0), sol=st.get("solucionado", 0),
                           by_crit=by_crit, by_area=by_area, amax=max([x[3] for x in by_area] + [1]), old=old, top=top, paths=paths)


@app.route("/findings")
@need()
def listing():
    lst, paths = locs()
    return render_template("list.html", items=findings(), locs=lst, paths=paths)


@app.route("/export.xlsx")
@need()
def export():
    _, paths = locs()
    wb = Workbook()
    ws = wb.active
    ws.title = "Hallazgos"
    ws.append(["ID", "Área", "TAG / equipo", "Zona", "Protección Ex", "Criticidad", "Estado", "Descripción",
               "Inspector", "Fecha inspección", "Resolución", "Cerrado por", "Fecha cierre", "Fotos antes", "Fotos después"])
    cnt = {}
    for r in rows("SELECT finding_id, kind, COUNT(*) n FROM photos GROUP BY 1, 2"):
        cnt[(r["finding_id"], r["kind"])] = r["n"]
    for f in findings():
        ws.append([f["id"], paths[f["location_id"]], f["tag"], f["zone"], f["protection"], f["criticality"],
                   f["status"], f["description"], f["inspector"], f["inspected_at"], f["resolution"], f["closer"],
                   f["closed_at"], cnt.get((f["id"], "antes"), 0), cnt.get((f["id"], "despues"), 0)])
    for col, wd in zip("ABCDEFGHIJKLMNO", (6, 45, 18, 7, 12, 10, 12, 60, 20, 14, 50, 20, 13, 11, 13)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="hallazgos_%s.xlsx" % date.today(),
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def equipment_list():
    return rows("SELECT e.*, l.name area FROM equipment e LEFT JOIN locations l ON l.id=e.location_id ORDER BY lower(e.tag)")


@app.route("/findings/new", methods=["GET", "POST"])
@need()
def new():
    lst, _ = locs()
    f = request.form
    if request.method == "POST":
        eq = None
        if f.get("equipment_id", "").isdigit():
            eq = db().execute("SELECT * FROM equipment WHERE id=%s", (int(f["equipment_id"]),)).fetchone()
        loc = f.get("location_id") or (str(eq["location_id"]) if eq and eq["location_id"] else "")
        zone = f.get("zone") or (eq["zone"] if eq else None)
        if not (loc.isdigit() and zone in ZONES and f.get("criticality") in CRIT
                and f.get("description", "").strip()):
            flash("Completá área, zona, criticidad y descripción.")
        else:
            r = run("INSERT INTO findings(location_id,tag,equipment_id,zone,protection,criticality,description,inspector_id,inspector_name,inspected_at)"
                    " VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s) RETURNING id", int(loc), eq["tag"] if eq else "", eq["id"] if eq else None,
                    zone, eq["protection"] if eq else None, f["criticality"],
                    f["description"].strip(), session["uid"], session["name"], dt(f.get("inspected_at"))).fetchone()
            add_photos(r["id"], "antes")
            return redirect(url_for("detail", fid=r["id"]))
    return render_template("form.html", locs=lst, equipment=equipment_list(), f=f, today=date.today().isoformat())


@app.route("/findings/<int:fid>")
@need()
def detail(fid):
    f = one("SELECT f.*, COALESCE(u.name, f.inspector_name, 'Usuario eliminado') inspector, "
            "CASE WHEN f.closed_at IS NOT NULL THEN COALESCE(c.name, f.closer_name, 'Usuario eliminado') END closer FROM findings f "
            "LEFT JOIN users u ON u.id=f.inspector_id LEFT JOIN users c ON c.id=f.closed_by WHERE f.id=%s", fid)
    ph = {"antes": [], "despues": []}
    for p in rows("SELECT id, kind FROM photos WHERE finding_id=%s ORDER BY id", fid):
        ph[p["kind"]].append(p)
    return render_template("detail.html", f=f, ph=ph, path=locs()[1][f["location_id"]], today=date.today().isoformat())


@app.post("/findings/<int:fid>/photos")
@need()
def photos(fid):
    one("SELECT 1 FROM findings WHERE id=%s", fid)
    kind = request.form.get("kind")
    if kind not in ("antes", "despues"):
        abort(403)
    if not add_photos(fid, kind):
        flash("No se agregó ninguna foto.")
    return redirect(url_for("detail", fid=fid))


@app.post("/findings/<int:fid>/close")
@need()
def close(fid):
    one("SELECT 1 FROM findings WHERE id=%s", fid)
    res = request.form.get("resolution", "").strip()
    if not res:
        flash("Describí qué se reparó o cambió para poder cerrar el hallazgo.")
    else:
        run("UPDATE findings SET status='solucionado', resolution=%s, closed_by=%s, closer_name=%s, closed_at=%s WHERE id=%s",
            res, session["uid"], session["name"], dt(request.form.get("closed_at")), fid)
        add_photos(fid, "despues")
    return redirect(url_for("detail", fid=fid))


@app.post("/findings/<int:fid>/reopen")
@need()
def reopen(fid):
    one("SELECT 1 FROM findings WHERE id=%s", fid)
    run("UPDATE findings SET status='pendiente', resolution=NULL, closed_by=NULL, closed_at=NULL WHERE id=%s", fid)
    return redirect(url_for("detail", fid=fid))


@app.route("/files/<int:pid>")
@need()
def files(pid):
    r = one("SELECT data FROM photos WHERE id=%s", pid)
    resp = send_file(io.BytesIO(r["data"]), mimetype="image/jpeg")
    resp.headers["Cache-Control"] = "private, max-age=86400"
    return resp


def eq_fields(f):
    tag, loc, zone, prot = f.get("tag", "").strip(), f.get("location_id", ""), f.get("zone", ""), f.get("protection")
    if not tag or (loc and not loc.isdigit()) or (zone and zone not in ZONES):
        return None
    return tag, f.get("description", "").strip(), int(loc) if loc else None, zone or None, prot if prot in PROT else None


@app.route("/equipment", methods=["GET", "POST"])
@need("supervisor")
def equipment():
    if request.method == "POST":
        v = eq_fields(request.form)
        if not v:
            flash("Escribí el tag.")
        else:
            try:
                run("INSERT INTO equipment(tag,description,location_id,zone,protection) VALUES(%s,%s,%s,%s,%s)", *v)
                flash("Tag agregado.")
            except psycopg.errors.UniqueViolation:
                flash("Ya existe un tag con ese nombre.")
            except psycopg.errors.ForeignKeyViolation:
                flash("El área elegida no existe.")
        return redirect(url_for("equipment"))
    q = request.args.get("q", "").strip()
    items = rows("SELECT e.*, l.name area, (SELECT COUNT(*) FROM findings f WHERE f.equipment_id=e.id) nf FROM equipment e "
                 "LEFT JOIN locations l ON l.id=e.location_id WHERE (%s='' OR e.tag ILIKE %s OR e.description ILIKE %s OR l.name ILIKE %s) "
                 "ORDER BY lower(e.tag)", q, "%" + q + "%", "%" + q + "%", "%" + q + "%")
    return render_template("equipment.html", items=items, areas=locs()[0], q=q)


@app.route("/equipment/<int:eid>/edit", methods=["GET", "POST"])
@need("supervisor")
def equipment_edit(eid):
    e = one("SELECT * FROM equipment WHERE id=%s", eid)
    if request.method == "POST":
        v = eq_fields(request.form)
        if not v:
            flash("Escribí el tag.")
        else:
            try:
                run("UPDATE equipment SET tag=%s, description=%s, location_id=%s, zone=%s, protection=%s WHERE id=%s", *v, eid)
                run("UPDATE findings SET tag=%s WHERE equipment_id=%s", v[0], eid)
                flash("Tag actualizado.")
                return redirect(url_for("equipment"))
            except psycopg.errors.UniqueViolation:
                flash("Ya existe otro tag con ese nombre.")
            except psycopg.errors.ForeignKeyViolation:
                flash("El área elegida no existe.")
    return render_template("equipment_edit.html", e=e, areas=locs()[0])


@app.post("/equipment/<int:eid>/delete")
@need("supervisor")
def equipment_delete(eid):
    one("SELECT 1 FROM equipment WHERE id=%s", eid)
    run("DELETE FROM equipment WHERE id=%s", eid)
    flash("Tag eliminado. Los hallazgos ya cargados conservan su texto.")
    return redirect(url_for("equipment"))


@app.route("/areas")
@need("admin")
def areas():
    return render_template("areas.html", items=rows(
        "SELECT l.id, l.name, (SELECT COUNT(*) FROM equipment e WHERE e.location_id=l.id) ne, "
        "(SELECT COUNT(*) FROM findings f WHERE f.location_id=l.id) nf, "
        "(SELECT COUNT(*) FROM findings f WHERE f.location_id=l.id AND f.status='pendiente') np FROM locations l "
        "ORDER BY array_position(%s::text[], l.name), l.name", AREAS))


@app.route("/users", methods=["GET", "POST"])
@need("admin")
def users():
    if request.method == "POST":
        f = request.form
        un, nm, role, pw = f.get("username", "").strip().lower(), f.get("name", "").strip(), f.get("role"), f.get("password", "")
        if not un or not nm or role not in ROLES or len(pw) < 8:
            flash("Usuario, nombre y rol son obligatorios; la contraseña debe tener al menos 8 caracteres.")
        else:
            try:
                run("INSERT INTO users(username,name,pw,role) VALUES(%s,%s,%s,%s)", un, nm, generate_password_hash(pw), role)
                flash("Usuario creado.")
            except psycopg.errors.UniqueViolation:
                flash("Ya existe un usuario con ese nombre de usuario.")
        return redirect(url_for("users"))
    return render_template("users.html", items=rows(
        "SELECT u.*, (SELECT COUNT(*) FROM findings f WHERE f.inspector_id=u.id) nf FROM users u ORDER BY lower(u.username)"))


@app.route("/users/<int:uid>/edit", methods=["GET", "POST"])
@need("admin")
def user_edit(uid):
    u = one("SELECT * FROM users WHERE id=%s", uid)
    if request.method == "POST":
        f = request.form
        un, nm, role, pw = f.get("username", "").strip().lower(), f.get("name", "").strip(), f.get("role"), f.get("password", "")
        active = bool(f.get("active"))
        if uid == session["uid"]:
            role, active = "admin", True
        if not un or not nm or role not in ROLES or (pw and len(pw) < 8):
            flash("Usuario, nombre y rol son obligatorios; la contraseña nueva debe tener al menos 8 caracteres.")
        else:
            try:
                run("UPDATE users SET username=%s, name=%s, role=%s, active=%s WHERE id=%s", un, nm, role, active, uid)
                run("UPDATE findings SET inspector_name=%s WHERE inspector_id=%s", nm, uid)
                run("UPDATE findings SET closer_name=%s WHERE closed_by=%s", nm, uid)
                if pw:
                    run("UPDATE users SET pw=%s WHERE id=%s", generate_password_hash(pw), uid)
                    run("DELETE FROM attempts WHERE username=%s", un)
                flash("Usuario actualizado.")
                return redirect(url_for("users"))
            except psycopg.errors.UniqueViolation:
                flash("Ya existe otro usuario con ese nombre de usuario.")
    return render_template("user_edit.html", u=u, me=(uid == session["uid"]))


@app.post("/users/<int:uid>/delete")
@need("admin")
def user_delete(uid):
    if uid == session["uid"]:
        flash("No podés eliminar tu propio usuario.")
        return redirect(url_for("users"))
    u = one("SELECT username FROM users WHERE id=%s", uid)
    run("DELETE FROM attempts WHERE username=%s", u["username"])
    run("DELETE FROM users WHERE id=%s", uid)
    flash("Usuario eliminado. Sus hallazgos se conservan con su nombre.")
    return redirect(url_for("users"))


@app.route("/password", methods=["GET", "POST"])
@need()
def password():
    if request.method == "POST":
        f = request.form
        u = one("SELECT username, pw FROM users WHERE id=%s", session["uid"])
        if one("SELECT COUNT(*) n FROM attempts WHERE username=%s AND at > now() - interval '10 minutes'", u["username"])["n"] >= 5:
            flash("Demasiados intentos fallidos. Esperá 10 minutos.")
        elif not check_password_hash(u["pw"], f.get("current", "")):
            run("INSERT INTO attempts(username) VALUES(%s)", u["username"])
            flash("La contraseña actual no es correcta.")
        elif len(f.get("new", "")) < 8:
            flash("La contraseña nueva debe tener al menos 8 caracteres.")
        elif f.get("new") != f.get("repeat"):
            flash("Las contraseñas nuevas no coinciden.")
        elif f.get("new") == f.get("current"):
            flash("La contraseña nueva debe ser distinta de la actual.")
        else:
            run("UPDATE users SET pw=%s WHERE id=%s", generate_password_hash(f["new"]), session["uid"])
            run("DELETE FROM attempts WHERE username=%s", u["username"])
            flash("Contraseña actualizada.")
            return redirect(url_for("dashboard"))
    return render_template("password.html")


# ---------- importación de tags desde planilla (.xlsx / .csv) ----------
SYN = {"tag": {"tag", "tags", "equipo", "tagequipo", "tagoequipo", "codigo"},
       "description": {"descripcion", "description", "desc", "detalle", "nombre"},
       "area": {"area"}, "zone": {"zona", "zone"},
       "protection": {"proteccion", "proteccionex", "protection", "tipodeproteccion", "tipoproteccion", "ex"}}
PROT_BARE = {"d", "e", "i", "p", "n", "m", "o", "q", "t", "db", "eb", "ia", "ib", "ic", "ma", "mb", "tb"}
PROT_NA = {"", "na", "n a", "n/a", "otro", "otro / n.a.", "otro n a", "no aplica", "ninguna", "-"}
MAX_ROWS = 5000


def _norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", s).strip()


def _key(s):
    return re.sub(r"[^a-z0-9]", "", _norm(s))


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def read_rows(f):
    name, raw = (f.filename or "").lower(), f.read()
    if name.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        out = [[_cell(c) for c in r] for r in wb.worksheets[0].iter_rows(values_only=True)]
        wb.close()
        return out
    if name.endswith((".csv", ".txt")):
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("cp1252")
        first = text.splitlines()[0] if text.strip() else ""
        d = max(";,\t", key=first.count)
        return [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), delimiter=d)]
    raise ValueError("Formato no admitido. Subí un archivo .xlsx o .csv (si es .xls, guardalo como .xlsx).")


def norm_prot(v):
    """Devuelve (valor, advertencia)."""
    n = _norm(v)
    if n in PROT_NA:
        return (None if not n else "Otro / N.A."), False
    m = re.match(r"^ex\s*([deipnmoqt])", n)
    if m:
        return "Ex " + m.group(1), False
    if n in PROT_BARE:
        return "Ex " + n[0], False
    return "Otro / N.A.", True


def import_tags(f, save):
    rows_ = [(n, r) for n, r in enumerate(read_rows(f), 1) if any(r)]
    if not rows_:
        raise ValueError("La planilla está vacía.")
    cols = {"tag": 0, "description": 1, "area": 2, "zone": 3, "protection": 4}
    hdr = {}
    for ci, h in enumerate(rows_[0][1]):
        for name, keys in SYN.items():
            if _key(h) in keys and name not in hdr:
                hdr[name] = ci
    if "tag" in hdr:
        cols, rows_ = hdr, rows_[1:]
    if len(rows_) > MAX_ROWS:
        raise ValueError("La planilla tiene más de %d filas. Dividila en partes." % MAX_ROWS)
    amap = {_key(n): i for i, n, _ in locs()[0]}
    out, errs, warns, dups = {}, [], [], 0
    for n, r in rows_:
        g = lambda k: r[cols[k]].strip() if k in cols and cols[k] < len(r) else ""
        tag, desc, area, zone, prot = g("tag"), g("description"), g("area"), g("zone"), g("protection")
        if not tag:
            errs.append((n, "falta el tag"))
            continue
        loc = amap.get(_key(area)) if area else None
        if area and not loc:
            errs.append((n, "área «%s» no válida" % area))
            continue
        z = None
        if zone:
            z = ZMAP.get(re.sub(r"^(zona|zone|z)\s*", "", _norm(zone)).replace(" ", ""))
            if not z:
                errs.append((n, "zona «%s» no válida" % zone))
                continue
        p, w = norm_prot(prot)
        if w:
            warns.append((n, "protección «%s» no reconocida: se cargó como «Otro / N.A.»" % prot))
        k = tag.lower()
        if k in out:
            dups += 1
            o = out[k]
            o["tag"] = tag
            for key, val in (("loc", loc), ("zone", z), ("desc", desc), ("prot", p)):
                if val:
                    o[key] = val
        else:
            out[k] = dict(tag=tag, desc=desc, loc=loc, zone=z, prot=p)
    ex = {r["tag"].lower(): r for r in rows("SELECT * FROM equipment")}
    ins, upd, same = [], [], 0
    for k, o in out.items():
        e = ex.get(k)
        if not e:
            ins.append((o["tag"], o["desc"], o["loc"], o["zone"], o["prot"]))
            continue
        new = (o["tag"], o["desc"] or e["description"] or "", o["loc"] or e["location_id"], o["zone"] or e["zone"], o["prot"] or e["protection"])
        if new == (e["tag"], e["description"] or "", e["location_id"], e["zone"], e["protection"]):
            same += 1
        else:
            upd.append(new + (e["id"],))
    if save and (ins or upd):
        try:
            with db().transaction():
                cur = db().cursor()
                cur.executemany("INSERT INTO equipment(tag,description,location_id,zone,protection) VALUES(%s,%s,%s,%s,%s)", ins)
                cur.executemany("UPDATE equipment SET tag=%s, description=%s, location_id=%s, zone=%s, protection=%s WHERE id=%s", upd)
                cur.executemany("UPDATE findings SET tag=%s WHERE equipment_id=%s", [(u[0], u[5]) for u in upd])
        except psycopg.Error:
            app.logger.exception("import de tags")
            raise ValueError("No se pudo guardar y no se aplicó ningún cambio. Probá de nuevo.")
    return dict(saved=save, total=len(rows_), new=len(ins), updated=len(upd), same=same, dups=dups, errs=errs, warns=warns)


@app.route("/equipment/import", methods=["GET", "POST"])
@need("admin")
def equipment_import():
    res = None
    if request.method == "POST":
        f = request.files.get("file")
        if not f or not f.filename:
            flash("Elegí un archivo .xlsx o .csv.")
        else:
            try:
                res = import_tags(f, request.form.get("mode") == "import")
            except ValueError as e:
                flash(str(e))
            except Exception:
                app.logger.exception("lectura de planilla")
                flash("No se pudo leer el archivo. Verificá que sea un .xlsx o .csv válido.")
    return render_template("equipment_import.html", res=res)


@app.route("/equipment/template.xlsx")
@need("admin")
def equipment_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Tags"
    ws.append(["Tag", "Descripción", "Área", "Zona", "Protección Ex"])
    ws.append(["TB-12", "Tablero de fuerza", AREAS[0], "1", "Ex d"])
    for col, wd in zip("ABCDE", (16, 40, 16, 8, 14)):
        ws.column_dimensions[col].width = wd
    v = wb.create_sheet("Valores válidos")
    v.append(["Áreas", "Zonas", "Protección Ex"])
    for i in range(max(len(AREAS), len(ZONES), len(PROT))):
        v.append([AREAS[i] if i < len(AREAS) else None, ZONES[i] if i < len(ZONES) else None, PROT[i] if i < len(PROT) else None])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="planilla_tags.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ---------- desclasificación de áreas ----------
@app.route("/declass")
@need()
def declass():
    a, w, p = request.args, [], []
    if a.get("loc", "").isdigit():
        w.append("d.location_id=%s")
        p.append(int(a["loc"]))
    if a.get("q"):
        w.append("(d.reason ILIKE %s OR d.description ILIKE %s OR d.responsible ILIKE %s OR EXISTS "
                 "(SELECT 1 FROM declass_tags t WHERE t.declass_id=d.id AND t.tag ILIKE %s))")
        p += ["%" + a["q"] + "%"] * 4
    items = rows("SELECT d.*, l.name area, (SELECT string_agg(t.tag, ', ' ORDER BY t.tag) FROM declass_tags t WHERE t.declass_id=d.id) tags, "
                 "(SELECT COUNT(*) FROM declass_tags t WHERE t.declass_id=d.id) nt FROM declassifications d "
                 "JOIN locations l ON l.id=d.location_id" + (" WHERE " + " AND ".join(w) if w else "")
                 + " ORDER BY d.done_at DESC, d.id DESC", *p)
    return render_template("declass.html", items=items, locs=locs()[0])


@app.route("/declass/new", methods=["GET", "POST"])
@need()
def declass_new():
    f = request.form
    if request.method == "POST":
        ids = sorted({int(x) for x in f.getlist("equipment_id") if x.isdigit()})
        loc = f.get("location_id", "")
        reason, desc, resp = f.get("reason", "").strip(), f.get("description", "").strip(), f.get("responsible", "").strip()
        eqs = db().execute("SELECT id, tag, zone FROM equipment WHERE id = ANY(%s) AND location_id=%s",
                           (ids, int(loc))).fetchall() if loc.isdigit() and ids else []
        if not (loc.isdigit() and reason and desc and resp):
            flash("Completá área, motivo, trabajo realizado y responsable.")
        elif not eqs or len(eqs) != len(ids):
            flash("Elegí al menos un tag del área seleccionada.")
        else:
            with db().transaction():
                did = run("INSERT INTO declassifications(location_id,reason,description,done_at,responsible,created_by,set_safe) "
                          "VALUES(%s,%s,%s,%s,%s,%s,%s) RETURNING id", int(loc), reason, desc, dt(f.get("done_at")), resp,
                          session["uid"], bool(f.get("set_safe"))).fetchone()["id"]
                db().cursor().executemany("INSERT INTO declass_tags(declass_id,equipment_id,tag,prev_zone) VALUES(%s,%s,%s,%s)",
                                          [(did, e["id"], e["tag"], e["zone"]) for e in eqs])
                if f.get("set_safe"):
                    run("UPDATE equipment SET zone='Segura' WHERE id = ANY(%s)", [e["id"] for e in eqs])
            add_photos(did, "antes", "declass_photos", "declass_id", "photos_antes")
            add_photos(did, "despues", "declass_photos", "declass_id", "photos_despues")
            return redirect(url_for("declass_detail", did=did))
    return render_template("declass_form.html", locs=locs()[0], equipment=equipment_list(), f=f,
                           today=date.today().isoformat(), me=session["name"])


@app.route("/declass/<int:did>")
@need()
def declass_detail(did):
    d = one("SELECT d.*, l.name area FROM declassifications d JOIN locations l ON l.id=d.location_id WHERE d.id=%s", did)
    ph = {"antes": [], "despues": []}
    for p in rows("SELECT id, kind FROM declass_photos WHERE declass_id=%s ORDER BY id", did):
        ph[p["kind"]].append(p)
    return render_template("declass_detail.html", d=d, ph=ph,
                           tags=rows("SELECT * FROM declass_tags WHERE declass_id=%s ORDER BY tag", did))


@app.post("/declass/<int:did>/photos")
@need()
def declass_add_photos(did):
    one("SELECT 1 FROM declassifications WHERE id=%s", did)
    kind = request.form.get("kind")
    if kind not in ("antes", "despues"):
        abort(400)
    if not add_photos(did, kind, "declass_photos", "declass_id"):
        flash("No se agregó ninguna foto.")
    return redirect(url_for("declass_detail", did=did))


@app.route("/dfiles/<int:pid>")
@need()
def dfiles(pid):
    r = one("SELECT data FROM declass_photos WHERE id=%s", pid)
    resp = send_file(io.BytesIO(r["data"]), mimetype="image/jpeg")
    resp.headers["Cache-Control"] = "private, max-age=86400"
    return resp
